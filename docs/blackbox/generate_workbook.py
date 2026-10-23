"""Generate the manual black-box test specification; does not execute application tests.
Run from any directory: python /workspace/librarymanagement_ByAI/docs/blackbox/generate_workbook.py
Requires openpyxl. All outputs are written beside this file and at repository root.
"""
import csv
import json
import shutil
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CASES = []
HEADERS = ['ID', 'ID cũ', 'Yêu cầu', 'Tên test case', 'Ưu tiên', 'Kỹ thuật', 'Nguồn kỳ vọng', 'Tiền điều kiện', 'Dữ liệu đầu vào', 'Các bước thực hiện', 'HTTP mong đợi', 'Kết quả mong đợi', 'Hậu điều kiện / dọn dữ liệu', 'Trạng thái', 'Kết quả thực tế', 'Mã lỗi', 'Bằng chứng']
REG = 'OpenAPI: Register; decisions.md: tài khoản/email/mật khẩu'
BOOK = 'OpenAPI: BookWrite, /admin/books; decisions.md: tồn kho'
BORROW = 'OpenAPI: BorrowCreate, BorrowDecision; decisions.md: mượn/xét duyệt'
AUTH = 'OpenAPI: securitySchemes và security của operation; decisions.md: phân quyền/session'
PAGING = 'OpenAPI: query parameters page/pageSize/q/category/sort'
USER = 'OpenAPI: UserCreate/UserUpdate; decisions.md: quản trị người dùng'

def add(old, fr, name, pre, data, steps, http, expect, source, technique='Phân vùng tương đương', priority='P1', cleanup=None):
    CASES.append([f'BB2-{len(CASES)+1:03}', old, fr, name, priority, technique, source,
                  'BASE; ' + pre, data, steps, str(http), expect,
                  cleanup or 'Đọc lại đối tượng bằng API nếu có thay đổi. Giữ bằng chứng; dọn theo CLEANUP trong Hướng dẫn.',
                  'Chưa chạy', '', '', ''])

def validation(old, fr, field, values, base, endpoint, source, pre='GUEST', good=201):
    for value, valid in values:
        label=json.dumps(value, ensure_ascii=False)
        if valid:
            if field == 'password':
                expected = '201; tài khoản mới được tạo. Đăng nhập bằng chính mật khẩu đầu vào trả 200; JSON không trả password/hash.'
                verify = '4. POST /auth/login với tài khoản vừa tạo và mật khẩu đầu vào.'
            elif field == 'totalQuantity':
                expected = f'data.totalQuantity={value}; data.availableQuantity={value}; mã/tên/tác giả đúng BOOK.'
                verify = '4. GET /books/{ID mới} bằng A1, xác nhận totalQuantity và availableQuantity.'
            elif field == 'email':
                expected = 'data.email=null.' if value in ['<OMIT>', '', None] else 'data.email đúng địa chỉ đã gửi.'
                verify = '4. A1 GET /admin/users?q=username, xác nhận đúng ID mới.'
            else:
                expected = f'data.{field} đúng giá trị đã gửi; các trường khác đúng dữ liệu nền; không lộ password/hash.'
                verify = '4. A1 GET danh sách users hoặc GET /books/{ID mới} tùy đối tượng để xác nhận dữ liệu.'
        else:
            expected = 'success=false; error.code=VALIDATION_ERROR; không tạo/sửa đối tượng.'
            verify = '4. A1 tìm đúng username hoặc bookCode mới qua danh sách công khai phù hợp, xác nhận không có ID mới.'
        if endpoint.endswith('/auth/register'):
            verify = '4. A1 GET /admin/users?q={RUN}, duyệt mọi trang (marker nằm trong username hoặc fullName không bị thay); đối chiếu ID trước/sau.'
            if not valid:
                expected = 'success=false; error.code=VALIDATION_ERROR; không có ID tài khoản mới thuộc marker RUN.'
        if endpoint.endswith('/admin/books'):
            verify = ('4. GET /books/{ID mới} bằng A1; đối chiếu các trường và tổng/tồn.' if valid else '4. A1 GET /books?category=CAT{RUN}, duyệt mọi trang; danh sách vẫn rỗng, không có ID sách mới.')
        if endpoint.endswith('/borrow-requests'):
            expected = 'success=false; error.code=VALIDATION_ERROR; lịch sử U1 không có yêu cầu mới; tồn B1 không đổi.'
            verify = '4. U1 GET /me/borrow-requests và GET /books/{B1}; đối chiếu lịch sử/tồn trước khi gửi.'
        if endpoint.endswith('/auth/login'):
            expected = 'success=false; error.code=VALIDATION_ERROR; GET /me bằng cùng jar trả 401.'
            verify = '4. GET /me bằng cùng jar.'
        add(old, fr, f'{field} = {label}', pre, f'{base}; thay duy nhất {field} bằng {label}',
            f'1. Chuẩn bị dữ liệu nền mới.\n2. Gửi {endpoint}.\n3. Đọc phản hồi.\n{verify}',
            good if valid else 400, expected, source,
            'Giá trị biên' if isinstance(value, (int,float)) or str(value).startswith('repeat(') else 'Phân vùng tương đương')


add('BB-001','FR-01','Đăng ký hợp lệ','GUEST; username/email mới','REG','1. POST /auth/register với REG.\n2. GET /admin/users bằng A1, q=username.','201','success=true; data.username/fullName/email đúng; data.role=USER; không có password/hash.',REG)
for field in ['username','fullName','password']:
    add('BB-002','FR-01',f'Thiếu {field}','GUEST',f'REG bỏ khóa {field}','1. POST /auth/register.\n2. A1 GET /admin/users?q={RUN}, duyệt mọi trang; không có ID mới thuộc marker RUN.','400','VALIDATION_ERROR; không tạo tài khoản.',REG)
for field, lens, maximum, old in [('username',[0,1,2,63,64,65],64,'BB-003'),('password',[7,8,9,127,128,129],128,'BB-005'),('fullName',[0,1,2,149,150,151],150,'BB-006')]:
    validation(old,'FR-01',field,[(f'repeat(a,{n})', 1<=n<=maximum if field!='password' else 8<=n<=maximum) for n in lens], 'REG mới; repeat(a,n) là chuỗi n ký tự a, không phải chuỗi biểu thức', 'POST /auth/register', REG, 'GUEST; username một ký tự cần bảo đảm chưa tồn tại; cô lập từng ca để tránh trùng')
for value, valid in [('abc def',False),('abc@def',False),('a-Z_0.9-',True)]:
    validation('BB-004','FR-01','username',[(value,valid)],'REG','POST /auth/register',REG,'GUEST; username thử chưa tồn tại')
for field,old in [('username','BB-007'),('email','BB-009')]:
    add(old,'FR-01',f'Đăng ký trùng {field}','GUEST; U1 có email hợp lệ',f'REG mới nhưng {field} trùng U1', '1. POST /auth/register.\n2. A1 tìm username mới và đọc U1.','409','error.code=DUPLICATE; không tạo tài khoản mới; U1 không thay đổi.',REG)
for value in ['<OMIT>','',None,'bb{RUN}@example.com','abc@']:
    validation('BB-008','FR-01','email',[(value,value!='abc@')],'REG; <OMIT> nghĩa là bỏ khóa','POST /auth/register',REG)
add('BB-010','FR-01','Chèn role ADMIN khi đăng ký','GUEST','REG thêm role=ADMIN','1. POST /auth/register.\n2. A1 kiểm tra username mới.','400','VALIDATION_ERROR; không tạo tài khoản với bất kỳ role nào.',REG,priority='P0')
for role in ['USER','ADMIN']:
    add('BB-011','FR-02',f'Đăng nhập {role}',f'GUEST; tài khoản {role} có mật khẩu thử đã biết',f'LOGIN của {role}', '1. POST /auth/login bằng cookie jar mới.\n2. Kiểm tra Set-Cookie.\n3. GET /me bằng cookie vừa nhận.','200; GET /me: 200',f'success=true; data.user đúng tài khoản {role}; cookie library_session có HttpOnly, SameSite=Strict; /me đúng ID/role, không lộ password/hash.',AUTH,priority='P0')
for label,data in [('Sai mật khẩu','LOGIN U1 nhưng password sai'),('Tài khoản không tồn tại','LOGIN username chữ/số xác nhận không tồn tại')]:
    add('BB-012','FR-02',label,'GUEST',data,'1. POST /auth/login.\n2. GET /me bằng jar đó.','401; GET /me: 401','INVALID_CREDENTIALS; không cấp phiên có thể dùng để gọi /me.',AUTH)
for field in ['username','password']:
    for value in ['', '<OMIT>']:
        validation('BB-013','FR-02',field,[(value,False)],'LOGIN; <OMIT> nghĩa là bỏ khóa','POST /auth/login','OpenAPI: Login')
add('BB-014','FR-02','Ngưỡng rate limit auth','Instance test riêng đã khởi động; chưa có request /auth từ IP thử; không có runner khác dùng IP này','101 POST /auth/login cùng IP; LOGIN sai mật khẩu, đúng cấu trúc','1. Gửi tuần tự 101 lần trong dưới 15 phút.\n2. Ghi HTTP của từng lần.','Lần 1–100: 401; lần 101: 429','100 lần đầu INVALID_CREDENTIALS; lần 101 RATE_LIMIT; /me với jar rỗng vẫn 401. Không chạy ca này chung instance/IP với ca auth khác.', 'decisions.md: 100 yêu cầu/15 phút/IP trên nhóm auth','Giá trị biên')

for path in ['/books','/books/{B1}']:
    add('BB-015','FR-03' if path=='/books' else 'FR-04','Guest bị chặn '+path,'GUEST; B1 tồn tại',f'GET {path}','1. Gửi GET không có cookie.','401','UNAUTHENTICATED; không có dữ liệu sách.',AUTH,'Bảng quyết định','P0')
for field,q in [('title','Sach{RUN}'),('author','Tacgia{RUN}'),('bookCode','CODE{RUN}')]:
    add('BB-016','FR-03','Tìm theo '+field,'U1; A1 tạo B1 với title=Sach{RUN} Alpha, author=Tacgia{RUN}, bookCode=CODE{RUN}',f'GET /books?q={q}','1. Gửi truy vấn.\n2. Duyệt mọi trang nếu có nhiều trang.','200',f'Có ID B1; mọi kết quả có chuỗi tìm ở title, author hoặc bookCode (không phân biệt hoa/dấu).', 'decisions.md: tìm chuỗi con; '+PAGING)
for q in ['THU VIEN {RUN}','thu vien {RUN}']:
    add('BB-017','FR-03','Tìm tiếng Việt: '+q,'U1; A1 tạo B1 title=Thư viện {RUN}',f'GET /books?q={q}','1. Gửi truy vấn URL-encoded.','200','Có B1 trong kết quả; không bỏ sót vì hoa thường hoặc dấu.','decisions.md: tìm không phân biệt hoa/dấu')
add('BB-018','FR-03','Tìm không có kết quả','U1; A1 xác nhận không có sách khớp chuỗi chữ/số này','GET /books?q=ZZZ{RUN}NoMatch999','1. Gửi truy vấn.','200','data.items=[]; data.total=0; data.totalPages=0.',PAGING)
add('BB-019','FR-03','Kết hợp tìm và thể loại','U1; A1 tạo B1/B2 có cùng token Sach{RUN}, category lần lượt CAT{RUN}A/CAT{RUN}B','GET /books?q=Sach{RUN}&category=CAT{RUN}A','1. Gửi truy vấn bằng U1.','200','Có B1, không có B2; mọi item vừa khớp q vừa có category=CAT{RUN}A.',PAGING)
for sort,ex in [('title','Các title trong tập dữ liệu Aaa{RUN}, Bbb{RUN}, Ccc{RUN} theo thứ tự tăng A/B/C.'),('available','availableQuantity theo thứ tự 3,2,1; không ràng buộc thứ tự khi bằng nhau.')]:
    add('BB-020','FR-03','Sắp xếp '+sort,'U1; A1 tạo 3 sách category=SORT{RUN}, title Aaa/Bbb/Ccc{RUN}, số khả dụng lần lượt 1/3/2',f'GET /books?category=SORT{{RUN}}&sort={sort}', '1. Gửi truy vấn.\n2. So sánh cả danh sách 3 item.','200',ex,'UI: Tên sách A → Z / Còn nhiều nhất; '+PAGING)
for field,values in [('page',[0,1,2]),('pageSize',[0,1,2,99,100,101])]:
    for value in values:
        ok=value>=1 and (field!='pageSize' or value<=100)
        add('BB-021','FR-03',f'Biên {field}={value}','U1; có ít nhất 3 sách',f'GET /books?{field}={value}','1. Gửi truy vấn.','200' if ok else '400',f'data.{field}={value}; items không vượt pageSize; các trường total/totalPages nhất quán.' if ok else 'VALIDATION_ERROR; không trả danh sách thành công.',PAGING,'Giá trị biên')
for field,lens in [('q',[254,255,256]),('category',[99,100,101])]:
    for n in lens:
        ok=n<=(255 if field=='q' else 100)
        add('BB-022','FR-03',f'{field} dài {n}','U1',f'GET /books?{field}=repeat(a,{n})','1. Tạo chuỗi n ký tự a, URL-encode.\n2. Gửi truy vấn.','200' if ok else '400','Trả danh sách (có thể rỗng), không lỗi validation.' if ok else 'VALIDATION_ERROR.',PAGING,'Giá trị biên')
add('BB-022','FR-03','Sort ngoài allowlist','U1','GET /books?sort=invalid','1. Gửi truy vấn.','400','VALIDATION_ERROR.',PAGING)
for label,query,ex in [('Trang đầu','page=1&pageSize=2','2 item; 2 ID đầu trong thứ tự đã chọn.'),('Trang giữa','page=2&pageSize=2','2 item tiếp theo; không trùng ID trang đầu.'),('Trang cuối','page=3&pageSize=2','1 item còn lại; hợp ba trang là đúng 5 ID, không trùng/mất.'),('Vượt trang cuối','page=4&pageSize=2','items=[]; total vẫn 5, totalPages vẫn 3.')]:
    add('Mới','FR-03',label,'U1; A1 tạo chính xác 5 sách trong category=PAGE{RUN}, title A/B/C/D/E{RUN}; không ghi dữ liệu đồng thời',f'GET /books?category=PAGE{{RUN}}&sort=title&{query}', '1. Đọc trang 1–3 làm mốc khi cần.\n2. Gửi truy vấn mục tiêu.','200',ex+' data.total=5; data.totalPages=3.',PAGING)
add('Mới','FR-03','Phân trang mặc định','U1; ít nhất 13 sách','GET /books không gửi page/pageSize','1. Gửi truy vấn.','200','data.page=1; data.pageSize=12; items có 12 phần tử.',PAGING)
for label,ident,http in [('Tồn tại','{B1}',200),('Không tồn tại','{MISSING_ID}',404),('ID bằng 0','0',400),('ID âm','-1',400),('ID chữ','abc',400)]:
    add('BB-023' if http!=400 else 'BB-024','FR-04',label,'U1; B1 tồn tại; MISSING_ID không tồn tại',f'GET /books/{ident}','1. Gửi GET.',''+str(http),'data.id=B1, đúng thông tin fixture.' if http==200 else ('BOOK_NOT_FOUND; không trả dữ liệu sách.' if http==404 else 'VALIDATION_ERROR.'),'OpenAPI: /books/{id}')

add('BB-025','FR-05','Tạo yêu cầu không giảm tồn','U1; B1 total=2, available=2; không có PENDING U1/B1','BORROW','1. POST /borrow-requests.\n2. GET /me/borrow-requests và GET /books/{B1}.','201; các GET: 200','data.id là ID mới; lịch sử có đúng ID, user.id=U1, status=PENDING, ngày D/D+1; availableQuantity vẫn 2.',BORROW,'Chuyển trạng thái','P0')
for offset in [-1,0,1]:
    add('BB-026','FR-05',f'Ngày trả D{offset:+}','U1; B1 mới, không có PENDING','BORROW, dueDate='+('D' if offset==0 else f'D{offset:+}'),'1. POST /borrow-requests.\n2. Đọc lịch sử U1 và tồn B1.','201' if offset==1 else '400','Có yêu cầu mới PENDING, tồn không đổi.' if offset==1 else 'INVALID_DATES; không tạo yêu cầu, tồn không đổi.',BORROW,'Giá trị biên','P0')
for field in ['borrowDate','dueDate']:
    for value in ['2027-02-29','2028-02-30','31/12/2028','']:
        validation('BB-027','FR-05',field,[(value,False)],'BORROW; trường ngày còn lại hợp lệ (2028-01-01 hoặc 2029-01-01 để bảo đảm thứ tự nếu ngày hợp lệ)','POST /borrow-requests',BORROW,'U1; B1 mới, không có PENDING')
add('BB-028','FR-05','Ngày nhuận hợp lệ','U1; B1 mới, không có PENDING','BORROW với borrowDate=2028-02-29, dueDate=2028-03-01','1. POST /borrow-requests.\n2. GET lịch sử bằng cùng U1.','201; GET: 200','Ngày đọc lại đúng 2028-02-29/2028-03-01, status=PENDING.',BORROW)
add('BB-029','FR-05','Mượn sách không tồn tại','U1; MISSING_ID không tồn tại','BORROW bookId=MISSING_ID','1. POST /borrow-requests.\n2. Đọc lịch sử U1.','404','BOOK_NOT_FOUND; không tạo yêu cầu.',BORROW)
for field,value in [('userId','{U2}'),('status','APPROVED'),('reviewedBy','{A1}')]:
    add('BB-030','FR-05','Chèn '+field,'U1; B1 mới, không có PENDING',f'BORROW thêm {field}={value}','1. POST /borrow-requests.\n2. Đọc lịch sử U1/U2.','400','VALIDATION_ERROR; không có yêu cầu mới hoặc thay đổi tồn.',BORROW,priority='P0')
add('BB-031','FR-05','Trùng PENDING','U1; đã có đúng một R1 PENDING của U1/B1','BORROW cùng B1','1. POST /borrow-requests.\n2. Đọc đủ các trang lịch sử U1.','409','DUPLICATE_PENDING; vẫn đúng một PENDING U1/B1; tồn không đổi.',BORROW,'Chuyển trạng thái','P0')
add('BB-032','FR-05','ADMIN không được tạo yêu cầu','A1; B1 tồn tại','BORROW','1. POST /borrow-requests bằng A1.','403','FORBIDDEN; không tạo yêu cầu.',AUTH,'Bảng quyết định','P0')
for state in ['APPROVED','DENIED']:
    add('BB-054','FR-05/FR-11',f'Gửi mới sau {state}',f'U1; yêu cầu cũ U1/B1 đã {state}; không còn PENDING U1/B1','BORROW','1. POST /borrow-requests.\n2. GET lịch sử.','201','Có ID mới PENDING; yêu cầu cũ giữ trạng thái; tồn trước/sau tạo mới không đổi.',BORROW,'Chuyển trạng thái')
for who in ['U1','U2']:
    add('BB-033','FR-06','Hồ sơ '+who,'U1 và U2 có phiên riêng',f'GET /me bằng {who}','1. Gửi GET với jar chỉ của tài khoản mục tiêu.','200',f'data.id/username/role đúng {who}; không có password, passwordHash hoặc token trong JSON.',AUTH)
add('BB-034','FR-07','Cách ly lịch sử','U1/U2 có lịch sử riêng và cookie riêng','GET /me/borrow-requests bằng U1','1. Gửi GET, đọc mọi trang.','200','Tất cả item.user.id=U1; không có ID yêu cầu của U2.',AUTH,priority='P0')
for state in ['PENDING','APPROVED','DENIED','INVALID']:
    add('BB-035','FR-07','Lọc lịch sử '+state,'U1 có ít nhất một yêu cầu mỗi trạng thái hợp lệ',f'GET /me/borrow-requests?status={state}','1. Gửi GET bằng U1.','400' if state=='INVALID' else '200','VALIDATION_ERROR.' if state=='INVALID' else f'items không rỗng; mọi item.status={state} và user.id=U1.','OpenAPI: /me/borrow-requests')
add('BB-036','FR-07','Lịch sử rỗng','USER mới U0 chưa có yêu cầu','GET /me/borrow-requests bằng U0','1. Gửi GET.','200','items=[]; total=0; totalPages=0.','OpenAPI: /me/borrow-requests')
add('BB-037','FR-08','Đăng xuất thu hồi phiên','U1 có phiên còn hạn; giữ bản cookie chỉ trong HTTP client cục bộ','POST /auth/logout; GET /me với chính cookie cũ','1. POST logout.\n2. Gửi lại GET /me kèm cookie cũ, không dựa vào việc browser đã xóa cookie.','204; GET: 401','Logout không có body; cookie cũ không được chấp nhận (UNAUTHENTICATED).',AUTH,'Chuyển trạng thái','P0')

admin_ops=[('GET','/admin/stats','FR-09','Không body'),('GET','/admin/books','FR-10','Không body'),('POST','/admin/books','FR-10','BOOK mới'),('PUT','/admin/books/{B1}','FR-10','BOOK giữ mã B1; title mới'),('DELETE','/admin/books/{B1}','FR-10','Không body; B1 chưa có lịch sử'),('GET','/admin/borrow-requests','FR-11','Không body'),('PATCH','/admin/borrow-requests/{R1}','FR-11','decision=APPROVE; R1 PENDING còn tồn'),('GET','/admin/users','FR-12','Không body'),('POST','/admin/users','FR-12','REG mới; role=USER'),('PUT','/admin/users/{U2}','FR-12','USERPUT U2, fullName mới'),('DELETE','/admin/users/{U2}','FR-12','Không body; U2 chưa có lịch sử')]
for method,path,fr,data in admin_ops:
    for who,http in [('GUEST',401),('U1',403)]:
        add('BB-038',fr,f'{who} bị chặn {method} {path}',f'{who}; A1 chuẩn bị B1/R1/U2 phù hợp và ghi mốc đối tượng; Content-Type JSON, Origin hợp lệ',data,f'1. Gửi {method} {path} bằng vai trò mục tiêu.\n2. A1 đọc lại đối tượng nếu đây là mutation.',http,('UNAUTHENTICATED' if http==401 else 'FORBIDDEN')+'; không lộ dữ liệu quản trị; mutation không đổi đối tượng.',AUTH,'Bảng quyết định','P0')
add('BB-038','FR-09','ADMIN xem thống kê','A1','GET /admin/stats','1. Gửi GET.','200','success=true; có totalBooks,totalCopies,availableCopies,totalUsers,totalRequests,pending,approved,denied; các tổng không âm.','OpenAPI: /admin/stats; UI: Tổng quan')
add('BB-039','FR-09','Thống kê sau thêm sách','A1; không có người ghi đồng thời; lấy S0 từ GET stats','BOOK mới total=3, available mặc định','1. POST /admin/books.\n2. GET stats S1.','POST: 201; GET: 200','S1.totalBooks=S0+1; totalCopies=S0+3; availableCopies=S0+3; các chỉ số user/request/status không đổi.','decisions.md: tồn kho; UI: Tổng quan')
add('BB-039','FR-09','Thống kê sau tạo yêu cầu','A1/U1; B1 còn tồn; không trùng PENDING; không có ghi đồng thời; lấy S0','BORROW bằng U1','1. POST /borrow-requests.\n2. A1 GET stats S1.','POST: 201; GET: 200','totalRequests/pending tăng 1; approved/denied, totalBooks/totalCopies/availableCopies/totalUsers không đổi.','decisions.md: tạo yêu cầu không giảm tồn; UI: Tổng quan')
add('BB-039','FR-09','Thống kê sau duyệt','A1; R1 PENDING còn tồn; không ghi đồng thời; lấy S0','decision=APPROVE','1. PATCH R1.\n2. GET stats S1.','PATCH: 200; GET: 200','pending giảm 1; approved tăng 1; availableCopies giảm 1; totalRequests/denied/totalBooks/totalCopies/totalUsers không đổi.',BORROW)

add('BB-040','FR-10','Tạo sách mặc định tồn','A1','BOOK; totalQuantity=2; bỏ availableQuantity','1. POST /admin/books.\n2. GET /books/{ID mới}.','201; GET: 200','totalQuantity=2; availableQuantity=2; mã/tên/tác giả khớp BOOK.',BOOK)
add('BB-041','FR-10','Mã sách trùng','A1; B1 đã tồn tại','BOOK có bookCode trùng B1','1. POST /admin/books.\n2. Tìm theo mã B1.','409','DUPLICATE; không tạo sách mới; B1 không thay đổi.',BOOK)
for field,lens,maxn in [('bookCode',[0,1,2,49,50,51],50),('title',[0,1,2,254,255,256],255),('author',[0,1,2,254,255,256],255)]:
    validation('BB-042','FR-10',field,[(f'repeat(a,{n})',1<=n<=maxn) for n in lens],'BOOK mới; repeat(a,n) là n ký tự a','POST /admin/books',BOOK,'A1; cô lập ca một ký tự để không trùng mã')
validation('BB-043','FR-10','totalQuantity',[(n,n in [0,1,2]) for n in [-1,0,1,2,1.5]],'BOOK mới; bỏ availableQuantity','POST /admin/books',BOOK,'A1')
for available in [-1,0,1,2,3]:
    ok=0<=available<=2
    add('BB-044','FR-10',f'Tạo total=2, available={available}','A1',f'BOOK mới; totalQuantity=2; availableQuantity={available}','1. POST /admin/books.\n2. Đọc lại ID mới nếu có.','201' if ok else '400',f'totalQuantity=2; availableQuantity={available}.' if ok else ('VALIDATION_ERROR; không tạo sách.' if available<0 else 'INVALID_QUANTITY; không tạo sách.'),BOOK,'Bảng quyết định')
for total in [1,2,4]:
    add('BB-045','FR-10',f'Sửa tổng thành {total} khi đang mượn 2','A1; B1 total=3, available=1; tạo 2 khoản APPROVED qua U1/U2; fixture mới cho mỗi ca',f'BOOK đầy đủ giữ bookCode/title/author B1; totalQuantity={total}; bỏ availableQuantity','1. PUT /admin/books/{B1}.\n2. GET /books/{B1}; đọc hai yêu cầu đã duyệt.','409' if total==1 else '200', 'QUANTITY_CONFLICT; B1 giữ total=3, available=1.' if total==1 else f'B1 total={total}, available={total-2}; hai yêu cầu vẫn APPROVED.',BOOK,'Giá trị biên','P0')
add('Mới','FR-10','Từ chối ghi đè tồn khả dụng','A1; B1 total=3, available=1 với 2 khoản APPROVED','BOOK B1 totalQuantity=4, availableQuantity=4 (đúng phải là 2)','1. PUT /admin/books/{B1}.\n2. GET sách.','409','QUANTITY_CONFLICT; B1 vẫn total=3, available=1.',BOOK,priority='P0')
for history in [False,True]:
    add('BB-046','FR-10','Xóa sách '+('có lịch sử' if history else 'chưa có lịch sử'), 'A1; B1 '+('có yêu cầu DENIED' if history else 'không có yêu cầu nào'), 'DELETE /admin/books/{B1}', '1. DELETE B1.\n2. GET B1; nếu có lịch sử, đọc yêu cầu cũ.', '409; GET: 200' if history else '204; GET: 404', 'HISTORY_CONFLICT; sách và lịch sử giữ nguyên.' if history else 'DELETE không body; GET BOOK_NOT_FOUND.',BOOK)

for decision,available,old in [('APPROVE',1,'BB-047'),('DENY',1,'BB-048'),('APPROVE',0,'BB-049'),('DENY',0,'Mới')]:
    success=decision=='DENY' or available>0
    final='APPROVED' if decision=='APPROVE' else 'DENIED'
    add(old,'FR-11',f'{decision} khi available={available}',f'A1; R1 PENDING của U1/B1, B1.available={available}; có thể tạo PENDING trước khi USER khác mượn cuốn cuối',f'{{"decision":"{decision}","note":"Ghi chu {decision}"}}', '1. PATCH /admin/borrow-requests/{R1}.\n2. GET danh sách yêu cầu bằng A1, lọc ID R1.\n3. GET sách B1 bằng U1.', '200' if success else '409', f'data.id=R1, status={final}, reviewedBy=A1; lịch sử có reviewedAt trong khoảng gửi/nhận, decisionNote đúng; available={available-(decision=="APPROVE")}.' if success else 'OUT_OF_STOCK; R1 vẫn PENDING, reviewedBy/reviewedAt/decisionNote chưa có; available=0.',BORROW,'Bảng quyết định','P0')
for state in ['APPROVED','DENIED']:
    for decision in ['APPROVE','DENY']:
        add('BB-050','FR-11',f'{state} không chuyển qua {decision}',f'A1; R1 đã {state}; ghi trạng thái/metadata và tồn ban đầu',f'decision={decision}; note=khongduocdoi','1. PATCH R1.\n2. Đọc lại R1 và B1.','409','BORROW_REQUEST_ALREADY_PROCESSED; toàn bộ trạng thái/metadata xét duyệt và tồn giữ nguyên.',BORROW,'Chuyển trạng thái','P0')
for same,old in [(True,'BB-051'),(False,'BB-052')]:
    add(old,'FR-11','Duyệt đồng thời '+('cùng yêu cầu' if same else 'tranh cuốn cuối'), 'A1 có hai jar hợp lệ; '+('R1 PENDING; B1.available=2' if same else 'R1/R2 PENDING của U1/U2 cùng B1.available=1'), 'Hai PATCH decision=APPROVE; cùng R1' if same else 'Hai PATCH APPROVE, một R1 và một R2', '1. Dùng hai HTTP client đồng bộ bắt đầu bằng barrier, lưu giờ bắt đầu/kết thúc.\n2. Gửi song song (không tuần tự).\n3. Đọc cả phản hồi, lịch sử và sách. Nếu không có bằng chứng gửi song song, đánh dấu Blocked thay vì Pass.', 'Một 200; một 409', 'R1 APPROVED; tồn chỉ giảm 1 thành 1; không đổi quyết định lần hai.' if same else 'Chỉ một yêu cầu APPROVED; yêu cầu còn lại PENDING; available=0, không âm; không mất yêu cầu.',BORROW,'Chuyển trạng thái / Đồng thời','P0')
for n in [0,499,500,501]:
    add('BB-053','FR-11',f'Ghi chú dài {n}','A1; R1 PENDING mới, B1.available=1',f'decision=APPROVE; note=repeat(a,{n}) (chuỗi thực, không phải biểu thức)','1. PATCH R1.\n2. Đọc lại yêu cầu và sách.','200' if n<=500 else '400', ('R1 APPROVED; tồn giảm đúng 1; decisionNote rỗng/null.' if n==0 else f'R1 APPROVED; decisionNote đúng {n} ký tự a; tồn giảm đúng 1.') if n<=500 else 'VALIDATION_ERROR; R1 vẫn PENDING; tồn không đổi.',BORROW,'Giá trị biên')
add('BB-053','FR-11','Quyết định ngoài enum','A1; R1 PENDING còn tồn','decision=INVALID','1. PATCH R1; đọc lại R1 và sách.','400','VALIDATION_ERROR; không đổi trạng thái, metadata hoặc tồn.',BORROW)
add('Mới','FR-11','Xét duyệt ID không tồn tại','A1; MISSING_ID chưa tồn tại','PATCH /admin/borrow-requests/{MISSING_ID}; decision=APPROVE','1. Gửi PATCH hợp lệ.','404','REQUEST_NOT_FOUND; không ảnh hưởng yêu cầu/sách khác.',BORROW)
for state in ['PENDING','APPROVED','DENIED']:
    add('Mới','FR-11','Admin lọc '+state,'A1; có ít nhất một yêu cầu mỗi trạng thái',f'GET /admin/borrow-requests?status={state}','1. GET và duyệt đủ trang.','200',f'Mọi item.status={state}; tổng khớp số item qua mọi trang.','OpenAPI: /admin/borrow-requests')

add('BB-055','FR-12','Admin tạo USER','A1','REG mới; role=USER','1. POST /admin/users.\n2. GET /admin/users?q=username.','201; GET: 200','ID mới; đúng username/fullName/email; role=USER; danh sách chứa tài khoản; không password/hash.',USER)
add('BB-055','FR-12','Admin sửa họ tên','A1; U2 tồn tại','USERPUT U2; fullName=Ten Moi {RUN}; không gửi password','1. PUT /admin/users/{U2}.\n2. GET danh sách tìm U2.','200; GET: 200','fullName cập nhật; ID/username/role giữ nguyên.',USER)
add('BB-055','FR-12','Admin xóa USER không lịch sử','A1; U2 chưa có yêu cầu mượn, không phải ADMIN','DELETE /admin/users/{U2}','1. DELETE U2.\n2. Tìm username U2.\n3. Thử đăng nhập U2.','204; GET: 200; LOGIN: 401','DELETE không body; danh sách không có U2; không thể đăng nhập U2.',USER)
for op in ['POST','PUT']:
    for field in ['username','email']:
        add('BB-056','FR-12',f'{op} người dùng trùng {field}','A1; U1/U2 tồn tại, email khác nhau và không rỗng',('REG mới' if op=='POST' else 'USERPUT U2')+f'; thay {field} bằng giá trị U1',f'1. {op} '+('/admin/users.' if op=='POST' else '/admin/users/{U2}.')+'\n2. Đọc lại U1/U2 và tìm tài khoản mới nếu POST.','409','DUPLICATE; không tạo mới; U1/U2 không thay đổi.',USER)
add('BB-057','FR-12','Đổi mật khẩu thu hồi phiên','A1; U1 có phiên còn hạn; biết mật khẩu thử cũ/mới khác nhau','USERPUT U1 thêm password=P_NEW','1. PUT U1.\n2. GET /me dùng cookie cũ U1.\n3. LOGIN mật khẩu cũ bằng jar mới.\n4. LOGIN mật khẩu mới.','PUT: 200; /me cũ: 401; LOGIN cũ: 401; LOGIN mới: 200','Cookie cũ bị thu hồi; chỉ mật khẩu mới dùng được; ID/role U1 giữ nguyên.',USER,'Chuyển trạng thái','P0')
add('BB-058','FR-12','Đổi role thu hồi phiên','A1; U2 là USER có phiên còn hạn, không phải A1','USERPUT U2; role=ADMIN','1. PUT U2.\n2. GET /me dùng cookie cũ U2.\n3. LOGIN lại U2; GET /admin/stats.','PUT: 200; /me cũ: 401; LOGIN mới: 200; stats: 200','Phiên cũ bị thu hồi; phiên mới nhận role ADMIN và dùng được quyền ADMIN.',USER,'Chuyển trạng thái','P0')
for op in ['DELETE','PUT']:
    add('BB-059','FR-12', 'Không được tự '+('xóa' if op=='DELETE' else 'hạ quyền'),'A1; ít nhất hai ADMIN để kiểm riêng quy tắc tự bảo vệ','Không body' if op=='DELETE' else 'USERPUT A1; role=USER',f'1. {op} /admin/users/{{A1}} bằng A1.\n2. GET /me bằng A1.','409; GET: 200','ADMIN_PROTECTED; A1 vẫn tồn tại, role=ADMIN; phiên A1 còn dùng được.',USER,'Bảng quyết định','P0')
add('BB-060','FR-12','Không xóa USER có lịch sử','A1; U2 có một yêu cầu DENIED','DELETE /admin/users/{U2}','1. DELETE U2.\n2. GET danh sách và lịch sử của U2.','409; GET: 200','HISTORY_CONFLICT; U2 và lịch sử còn nguyên.',USER)
add('Mới','FR-12','Sửa hồ sơ không đổi mật khẩu','A1; U2 có mật khẩu thử đã biết và phiên còn hạn','USERPUT U2 đổi fullName; bỏ password và role','1. PUT U2.\n2. GET /me bằng phiên U2 hiện tại.\n3. LOGIN U2 bằng mật khẩu cũ trong jar khác.','PUT: 200; /me: 200; LOGIN: 200','Đổi đúng fullName; mật khẩu/role không đổi; phiên hiện tại không bị thu hồi.',USER)
for op,path,data,fr,code in [('PUT','/admin/books/{MISSING_ID}','BOOK','FR-10','BOOK_NOT_FOUND'),('DELETE','/admin/books/{MISSING_ID}','Không body','FR-10','BOOK_NOT_FOUND'),('PUT','/admin/users/{MISSING_ID}','REG; role=USER','FR-12','USER_NOT_FOUND'),('DELETE','/admin/users/{MISSING_ID}','Không body','FR-12','USER_NOT_FOUND')]:
    add('Mới',fr,op+' đối tượng không tồn tại','A1; MISSING_ID không tồn tại',data,f'1. Gửi {op} {path}.','404',f'{code}; không thay đổi dữ liệu khác.','OpenAPI: operation tương ứng')

add('BB-061','BR/Auth','Phiên hết hạn tuyệt đối','U1 có phiên mới; SESSION_HOURS=8; không đổi mật khẩu/role/logout; đồng hồ client/server đồng bộ','GET /me cùng cookie tại t0+7h59m và t0+8h01m; t0 là thời điểm nhận LOGIN','1. Giữ cookie trong client cục bộ.\n2. Gửi GET ở hai mốc, buộc gửi cùng cookie kể cả sau hạn cookie.\n3. Không đăng nhập lại; không sửa DB/đồng hồ.','Trước hạn: 200; sau hạn: 401','Trước hạn đúng U1; sau hạn UNAUTHENTICATED. Ghi giờ và độ lệch; ca chưa chờ đủ không được Pass.',AUTH,'Giá trị biên',cleanup='Ca kéo dài 8 giờ, chạy riêng. Không lưu cookie vào báo cáo/Git; xóa cookie cục bộ sau khi thu bằng chứng.')
add('BB-062','BR/Security','Origin khác nguồn bị chặn','U1; B1 còn tồn; không trùng PENDING; HTTP client gửi cookie chủ động, không dựa vào browser SameSite','BORROW; Origin=https://other.example','1. POST /borrow-requests kèm cookie U1.\n2. Đọc lịch sử U1.','403','ORIGIN_REJECTED; không có yêu cầu mới, tồn không đổi.','decisions.md: same-origin mutation',priority='P0')
add('BB-063','BR/Security','Chuỗi SQL được xử lý như dữ liệu','U1; có ít nhất một sách nhưng không có sách chứa chuỗi đầu vào','GET /books?q='+"' OR 1=1 --",'1. URL-encode chuỗi, GET /books.','200','items=[]; total=0; không trả toàn bộ danh sách và không lộ lỗi SQL.','decisions.md: LIKE chuỗi con, parameterized input')
add('BB-063','BR/Security','Không thực thi HTML trong tiêu đề','A1 tạo B1 riêng; U1 có browser mới','BOOK title=<script>window.bbXssFlag=1</script>','1. A1 POST tạo sách.\n2. U1 tìm mã B1, mở danh sách và chi tiết.\n3. Kiểm tra console window.bbXssFlag trên trang.','POST: 201; GET/UI: 200','Tiêu đề hiển thị dạng văn bản đúng nội dung; window.bbXssFlag là undefined; không tạo phần tử script từ dữ liệu tiêu đề.','traceability.md: NFR-SEC; UI hiển thị dữ liệu người dùng')
for width,height in [(1440,1060),(390,844)]:
    for role in ['USER','ADMIN']:
        add('BB-064','UI',f'Bố cục {role} {width}x{height}',f'Browser Chromium; {role} có tài khoản hợp lệ; B1 tồn tại',f'Viewport {width}x{height}', '1. Đăng nhập đúng vai trò.\n2. Mở danh sách, tìm B1, mở chi tiết.\n3. '+('Mở menu quản trị, màn hình quản lý sách.' if role=='ADMIN' else 'Mở Sách của tôi và Hồ sơ.')+'\n4. Kiểm tra cuộn ngang, nút và nội dung bị che.', 'N/A (UI)', 'document.documentElement.scrollWidth <= innerWidth; đọc được nội dung; điều khiển dùng được; dialog đóng được; không có pageerror mới.','README: responsive desktop/mobile')
add('BB-064','UI','Phản hồi form thiếu dữ liệu','GUEST; Chromium desktop','Form đăng ký để trống username, các trường còn lại hợp lệ','1. Mở Đăng ký.\n2. Điền các trường còn lại.\n3. Bấm Đăng ký.','N/A (UI)','Không tạo tài khoản; validation chỉ ra trường bắt buộc (native browser hoặc thông báo ứng dụng); vẫn ở form, không hiện lỗi kỹ thuật.','OpenAPI: Register.required; UI đăng ký')

# SRS original document review: add missing explicit obligations, keep old IDs stable.
for path,fr in [('/me','FR-06'),('/me/borrow-requests','FR-07')]:
    add('SRS-review',fr,'Guest bị chặn '+path,'GUEST',f'GET {path}','1. GET không kèm cookie.','401','UNAUTHENTICATED; không trả hồ sơ/lịch sử.','SRS FR-06/07; BR-03; HTTP conventions',priority='P0')
add('SRS-review','FR-05','Guest không tạo được yêu cầu','GUEST; B1 tồn tại','BORROW','1. POST /borrow-requests không cookie.\n2. A1 kiểm tra danh sách requests.','401','UNAUTHENTICATED; không tạo yêu cầu, tồn không đổi.','SRS FR-05; BR-11',priority='P0')
add('SRS-review','FR-06','Session không hợp lệ','Cookie library_session=invalidtoken; không phải cookie thật','GET /me','1. GET /me kèm cookie giả.','401','UNAUTHENTICATED; không lộ hồ sơ.','SRS FR-06; NFR-SEC-04',priority='P0')
add('SRS-review','FR-02','Đăng nhập API bằng email','U1 có email=bb{RUN}@example.com đã cấu hình','LOGIN username=email U1, password=P','1. POST /auth/login bằng jar mới.\n2. GET /me.','200; GET: 200','data.user.id=U1; /me đúng U1; có phiên hợp lệ.','SRS FR-02: username/email + password',priority='P0')
add('SRS-review','FR-02','Đăng nhập UI bằng email dài hơn 64','A1 tạo U1 có email=repeat(a,60)+@ui{RUN}.example.com; tổng 83 ký tự; browser Chromium; bắt đầu GUEST','Email 83 ký tự hợp lệ; password=P của U1','1. Xác nhận API LOGIN với email này trả 200 rồi logout.\n2. Mở form đăng nhập bằng browser.\n3. Nhập đầy đủ email bằng thao tác bàn phím; nhập password.\n4. Bấm Đăng nhập; quan sát giá trị ô, HTTP và màn hình.','API: 200; UI LOGIN: 200; /me: 200','Ô đăng nhập giữ đủ 83 ký tự; đăng nhập thành công, vào Khám phá sách; /me đúng U1. Không sửa kỳ vọng thành 401 theo lỗi hiện tại.','SRS FR-02; dictionary email VARCHAR(254); OpenAPI Login.username max254',priority='P0')
for method in ['POST','PUT']:
    add('SRS-review','FR-12',method+' role không hợp lệ','A1; U2 là USER còn tồn tại','REG mới role=INVALID' if method=='POST' else 'USERPUT U2 role=INVALID',f'1. {method} '+('/admin/users.' if method=='POST' else '/admin/users/{U2}.')+'\n2. Đọc lại danh sách/đối tượng.','400','VALIDATION_ERROR; không tạo mới; không đổi role/thuộc tính U2.','SRS FR-12 ngoại lệ role invalid; role USER/ADMIN')
add('SRS-review','FR-01','JSON không hợp lệ','GUEST','Content-Type application/json; raw body={"username":','1. POST /auth/register với raw body đúng như trên, không dùng JSON serializer tự sửa.','400','error.code=INVALID_JSON; success=false; không lộ stack/SQL.','SRS HTTP conventions; NFR-MNT-02; NFR-SEC-04/05')
add('SRS-review','FR-07','Phân trang lịch sử sai','U1','GET /me/borrow-requests?page=0','1. Gửi GET bằng U1.','400','VALIDATION_ERROR; không trả danh sách thành công.','SRS FR-07 ngoại lệ pagination invalid; OpenAPI page minimum1')

QUESTIONS = [
 ['Q-01','Mới cập nhật / sort=newest','UI ghi Mới cập nhật nhưng static review API cho thấy sắp theo ID tạo mới; OpenAPI chỉ nêu enum.','Cần chủ sản phẩm xác định theo thời điểm tạo hay thời điểm sửa, chiều và quy tắc bằng nhau. Không biến hành vi hiện tại thành chuẩn. Chưa có TC nghiệm thu sort=newest.'],
 ['Q-02','Query ngoài hợp đồng','userId không nằm trong query của /me/borrow-requests.','Bắt buộc không lộ dữ liệu U2; cần chốt HTTP 400 hoặc bỏ qua với 200. Chưa đưa biến thể này vào bộ TC thực thi cho tới khi chốt HTTP.'],
 ['Q-03','ADMIN cuối cùng','Quy tắc có trong decisions.md; không thể dựng tình huống một ADMIN hợp lệ hạ/xóa ADMIN cuối khác chỉ bằng luồng công khai mà không đồng thời vi phạm tự bảo vệ.','Đã có TC tự bảo vệ; không tuyên bố đã phủ riêng mọi nhánh last-admin. Cần thiết kế quyền nhiều ADMIN/fixture ngoài phạm vi hộp đen trước khi bổ sung.'],
 ['Q-04','Giới hạn triển khai ngoài SRS','Đã nhận SRS v1.0 ngày 16/10/2026. SRS dùng BIGINT UNSIGNED trong dictionary; code dùng INT và path max2147483647; totalQuantity max1.000.000 và description max10.000 chưa được chốt trong SRS.','Đây là khác biệt thiết kế/contract cần xác nhận, không tự gắn thành requirement đã được duyệt.'],
 ['Q-06','TBD của SRS so với decisions.md','SRS vẫn để TBD: auth/session/expiry, password, email, quyền Guest, Admin mượn, duplicate pending, quyền sửa tồn và tự quản lý ADMIN. decisions.md ghi các lựa chọn triển khai; SRS tải lên chưa cập nhật chúng.','TC của các chính sách này chỉ là kiểm tra contract/cấu hình hiện tại, chưa đủ để kết luận đạt SRS nếu thiếu phê duyệt thay đổi yêu cầu.'],
 ['Q-07','Mâu thuẫn quyền lịch sử Admin trong SRS','API catalog ghi /me/borrow-requests: User; authorization matrix ghi Admin được own borrow history. Code chỉ yêu cầu auth nên Admin cũng gọi được.','Cần chốt và cập nhật SRS; chưa kết luận code vi phạm ở mục này.'],
 ['Q-05','NFR và audit/rollback','Chưa có ngưỡng SLA/p95/RPO/RTO hoặc API đọc audit và cơ chế gây lỗi hạ tầng được hỗ trợ.','Không đánh dấu các yêu cầu này đạt. Bộ hiện tại không thay thế load test, penetration test, kiểm thử khôi phục hoặc kiểm thử hộp trắng rollback.']
]
GUIDE = [
 ['Nguồn SRS','Đặc tả yêu cầu phần mềm website quản lý thư viện — SRS.docx; v1.0 ngày 16/10/2026; SHA-256 a440fa65c73192e27e82b27dee6a8bbef78505d31ae98023388250db3356e703. Chỉ lưu metadata/trích dẫn mục; không tự đưa file người dùng đính kèm lên GitHub.'],
 ['Phiên bản','2.1 — đã đối chiếu SRS gốc do người dùng cung cấp và code tại acdb8a9. ID BB2-xxx mới; cột ID cũ giữ liên hệ BB-001…064. Số lượng tăng do tách đầu vào và thêm nhánh còn thiếu.'],
 ['Trạng thái','Các TC giữ Chưa chạy để không lẫn hồ sơ thực thi với mẫu thiết kế. Lượt review ngày 23/10/2026 đã chạy 19 tests có sẵn, browser workflow và 43 checks bổ trợ; có lỗi UI email xác nhận. Xem sheet Kiem chung và srs-code-review.md, không suy ra toàn bộ TC đã Pass.'],
 ['BASE','Môi trường test riêng, dữ liệu không phải production. Node/MySQL/app đã sẵn sàng; không có người ghi đồng thời trừ ca concurrency. MIN_PASSWORD_LENGTH=8, SESSION_HOURS=8, BLOCK_DUPLICATE_PENDING=true. Mỗi ca fixture mới, không phụ thuộc ca trước.'],
 ['API','Mọi path trong TC có prefix /api/v1; Content-Type application/json cho body. GUEST là jar rỗng. U1/U2 là hai USER khác nhau, A1 là ADMIN. A1 chuẩn bị dữ liệu, người thực hiện request mục tiêu theo tiền điều kiện.'],
 ['RUN','RUN là 8 ký tự chữ/số duy nhất mỗi ca. {ID} lấy từ phản hồi tạo fixture. repeat(a,n) là CHUỖI gồm đúng n ký tự a, không phải biểu thức gửi nguyên văn. <OMIT> nghĩa là bỏ khóa JSON. Ca chuỗi 1 ký tự phải chọn ký tự còn trống hoặc fixture cô lập.'],
 ['REG','JSON nền: username=bb{RUN}, fullName=Nguyen An {RUN}, email=bb{RUN}@example.com, password=P. P là mật khẩu thử hợp lệ dài 12 ký tự do tester tạo và giữ cục bộ. Biến thể password có thể thay bằng chuỗi a đúng độ dài n; không áp quy tắc chữ hoa/ký tự đặc biệt.'],
 ['LOGIN','JSON username của tài khoản mục tiêu, password=P đã dùng khi tạo tài khoản đó; mỗi tài khoản jar riêng. P_NEW khác P, dài 12 ký tự. Không đưa mật khẩu/cookie/token vào bằng chứng hoặc Git.'],
 ['BOOK','JSON nền: bookCode=CODE{RUN}, title=Sach {RUN}, author=Tacgia {RUN}, category=CAT{RUN}, totalQuantity=2; bỏ availableQuantity nếu TC không nói khác. Khi PUT giữ đủ bookCode,title,author,totalQuantity và giữ category/description hiện có để chỉ đổi biến mục tiêu.'],
 ['USERPUT','JSON giữ username, fullName, email hiện tại của tài khoản mục tiêu; bỏ password và role trừ khi TC yêu cầu thay. Tránh vô tình xóa email hoặc đổi trường không phải mục tiêu.'],
 ['BORROW','JSON nền: bookId=ID sách B1, borrowDate=D, dueDate=D+1. D là ngày thực hiện +1 ngày, dạng YYYY-MM-DD; cộng ngày theo lịch. Trường thiếu ID không được gửi dưới dạng chuỗi placeholder.'],
 ['Fixtures trạng thái','Tạo B1 mới qua A1; đăng ký/đăng nhập U1/U2; tạo request bằng USER rồi APPROVE/DENY bằng A1. Để có PENDING với tồn 0: tạo 2 request khi tồn 1, duyệt request USER còn lại. Không sửa DB để dựng fixture hộp đen.'],
 ['MISSING_ID','Chọn ID dương trong miền hợp đồng, xác nhận không có đối tượng mục tiêu qua GET danh sách đầy đủ hoặc GET chi tiết nếu được hỗ trợ; không suy luận từ một trang danh sách. Không có ghi đồng thời khi xác nhận.'],
 ['Chứng cứ','Mỗi TC ghi commit ứng dụng, cấu hình không bí mật, môi trường, browser, ngày giờ, input thực, HTTP/body đã che secret, mốc trước/sau. Lưu ID đối tượng và request ID; cookie không đưa vào báo cáo.'],
 ['Oracle chung','HTTP và nội dung/nghiệp vụ phải đồng thời đúng. Lỗi có success=false, error.code/message và requestId; không lộ SQL/stack/credential. Thành công JSON có success=true và data; 204 không body. Lỗi không làm thay đổi đối tượng nghiệp vụ.'],
 ['CLEANUP','Đối tượng tạo thành công chưa có lịch sử: A1 DELETE sách/user thử sau ca. Có lịch sử: chỉ giữ trong môi trường test riêng; muốn reset phải dùng cơ chế môi trường test được cho phép, không xóa volume/DB có dữ liệu cần giữ. Không để cleanup làm mất bằng chứng.'],
 ['Thứ tự chạy','Chuẩn bị fixture độc lập, chạy P0 trước P1. Rate limit dùng instance/IP riêng hoặc chạy cuối. Hết hạn session cần 8 giờ, đánh dấu Chưa chạy/Blocked nếu chưa chờ đủ.'],
 ['Phân loại kết quả','Pass: mọi assertion đạt. Fail: sai khác có bằng chứng. Blocked: thiếu fixture/công cụ/đặc tả. Chưa chạy: chưa thực hiện. Các mục chưa chốt nằm riêng ở sheet Cần làm rõ, chưa tính là TC thực thi.'],
 ['Phạm vi','Đã đọc SRS v1.0 ngày 16/10/2026 do người dùng cung cấp (SHA-256 trong srs-code-review.md). Phủ các nhóm FR-01…12 nhưng vẫn còn mục TBD/NFR chưa được nghiệm thu. Không gồm trả sách, gia hạn, phạt, thanh toán. Bộ này thiên về API hộp đen, có các ca UI riêng; không thay thế UAT toàn bộ màn hình.'],
 ['Sinh lại file','python docs/blackbox/generate_workbook.py (cần openpyxl). Script chỉ sinh CSV/Markdown/Excel, không chạy ứng dụng hoặc sửa mã ứng dụng.']
]

def style(ws, widths, height=100):
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for idx,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(idx)].width=width
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='1F4E78');cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height=height
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if row[0].row%2==0: cell.fill=PatternFill('solid',fgColor='EEF4FA')
    ws.sheet_view.zoomScale=80


CODE_REFS = {
 'FR-01':'server/app.js:37 (schemas), :186 (register); src/main.jsx:507 AuthForm',
 'FR-02':'server/app.js:168 (rate limit), :200 (login); src/main.jsx:507 AuthForm',
 'FR-03':'server/app.js:70 (paging), :251 listBooks; src/main.jsx:676 Catalog',
 'FR-04':'server/app.js:293; src/main.jsx:979 BookDetail',
 'FR-05':'server/app.js:300 (dates), :308 (create); src/main.jsx:1018 BorrowForm',
 'FR-06':'server/app.js:129 (auth), :250 (/me); src/main.jsx:1709 Profile',
 'FR-07':'server/app.js:348 listRequests, :396 (own); src/main.jsx:1191 Requests',
 'FR-08':'server/app.js:235 logout; src/main.jsx:195 logout',
 'FR-09':'server/app.js:400 stats; src/main.jsx:1374 Dashboard',
 'FR-10':'server/app.js:56 bookSchema, :413 create, :441 update, :481 delete; src/main.jsx:1102 BookForm',
 'FR-11':'server/app.js:348, :490, :491 (decisions); src/main.jsx:1191 Requests, :1318 DecisionForm',
 'FR-12':'server/app.js:539 list, :554 create, :568 update, :605 delete; src/main.jsx:1497 UserList, :1606 UserForm'
}
def alignment(row):
    old, fr, name = row[1], row[2], row[3]
    refs = '; '.join(CODE_REFS[k] for k in CODE_REFS if k in fr)
    basis = 'SRS '+fr+'; chi tiết mã lỗi/trường JSON theo OpenAPI. Đối chiếu tĩnh, không đồng nghĩa TC đã chạy.'
    extra = {
      'BB-004':'SRS không chốt regex username; theo OpenAPI hiện tại.',
      'BB-005':'SRS NFR-SEC-07 / TBD-SEC-01 vẫn TBD; 8–128 theo decisions.md/OpenAPI.',
      'BB-008':'SRS FR-01 / TBD-USER-01 vẫn TBD; email tùy chọn/định dạng theo decisions.md.',
      'BB-009':'SRS FR-01: duy nhất nếu email bật; cấu hình bật theo decisions.md.',
      'BB-011':'SRS FR-02; cookie session là lựa chọn bổ sung TBD-AUTH-01, không phải yêu cầu JWT bắt buộc.',
      'BB-014':'Rate-limit100/15 phút từ decisions.md; SRS không ấn định con số.',
      'BB-015':'SRS TBD-BOOK-01: quyền Guest chưa chốt; TC401 theo decisions.md.',
      'BB-017':'SRS checklist: search case/accent chưa chốt; TC theo decisions.md.',
      'BB-020':'SRS chỉ nêu sort; chiều A-Z/available theo nhãn UI và OpenAPI. newest tách Q-01.',
      'BB-021':'SRS nêu pagination invalid; ngưỡng1/100 từ decisions.md/OpenAPI, không phải con số SRS.',
      'BB-022':'Giới hạn query/enum từ OpenAPI; SRS nêu query phải hợp lệ.',
      'BB-031':'SRS BR-18 / TBD-BORROW-03 còn TBD; chặn trùng là decisions.md.',
      'BB-032':'SRS authorization matrix để TBD Admin mượn; TC403 theo decisions.md.',
      'BB-037':'SRS FR-08 phụ thuộc cơ chế auth; thu hồi ngay theo session trong decisions.md.',
      'BB-044':'SRS BR-10 quy định bất đẳng thức; quyền nhập available khi tạo vẫn TBD, TC theo decisions.md.',
      'BB-045':'SRS BR-10; quy tắc sửa tổng giữ số đang mượn là decisions.md.',
      'BB-057':'SRS không quy định cụ thể thu hồi session sau đổi password; theo OpenAPI/decisions.md.',
      'BB-058':'SRS FR-12; thu hồi session sau đổi role theo OpenAPI; quản trị ADMIN khác còn TBD.',
      'BB-059':'SRS FR-12 / TBD-ADMIN-01 còn TBD; tự bảo vệ ADMIN theo decisions.md.',
      'BB-061':'SRS NFR-SEC-08 / TBD-AUTH-02 còn TBD; 8 giờ theo decisions.md.',
      'BB-062':'SRS NFR-SEC; Origin check là cơ chế session/same-origin bổ sung trong decisions.md.',
      'BB-064':'Ca UI bổ sung; SRS chưa chốt viewport 1440/390 hoặc tiêu chí responsive chi tiết.'
    }
    if old in extra:basis=extra[old]
    if fr=='UI':refs='src/main.jsx:167 App, :648 Pagination, :676 Catalog; src/style.css'
    if fr=='BR/Auth':refs='server/app.js:129 auth, :217 session expiry'
    if fr=='BR/Security' and old=='BB-063':basis='SRS NFR-SEC-04/05 về input/error; ca SQL/HTML bổ trợ kiểm tra bảo mật web, không thay penetration test.'
    if fr=='BR/Security':refs='server/app.js:13 helmet, :150 Origin, :251 listBooks, :638 error handler; src/main.jsx:676 Catalog'
    if name=='Phân trang mặc định':basis='SRS checklist để pagination default TBD; default12 là decisions.md/OpenAPI.'
    return [basis, refs]

def main():
    ids=[r[0] for r in CASES]
    assert len(ids)==len(set(ids))
    assert all(len(r)==len(HEADERS) and all(r[i] for i in range(13)) for r in CASES)
    assert all(r[13]=='Chưa chạy' for r in CASES)
    output_headers=HEADERS+['Căn cứ SRS / phạm vi','Vị trí đối chiếu code (tĩnh)']
    output_cases=[r+alignment(r) for r in CASES]
    assert {f'BB-{i:03}' for i in range(1,65)} <= {r[1] for r in CASES}
    with (ROOT/'test-cases.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f, lineterminator="\n");writer.writerow(output_headers);writer.writerows(output_cases)
    wb=Workbook(); guide=wb.active;guide.title='Huong dan';guide.append(['Mục','Nội dung'])
    for r in GUIDE:guide.append(r)
    style(guide,[26,135],90)
    ws=wb.create_sheet('Test cases');ws.append(output_headers)
    for r in output_cases:ws.append(r)
    style(ws,[15,15,18,42,12,24,50,68,65,85,35,90,60,18,50,20,45,90,85],150)
    ws.freeze_panes='D2'
    val=DataValidation(type='list',formula1='"Chưa chạy,Pass,Fail,Blocked"');val.errorTitle='Trạng thái không hợp lệ';val.error='Chọn một trạng thái trong danh sách.';val.showErrorMessage=True;val.errorStyle='stop';ws.add_data_validation(val);val.add(f'N2:N{len(CASES)+1}')
    for row in ws.iter_rows(min_row=2):
        if row[4].value=='P0': row[4].font=Font(bold=True,color='C00000')
    trace=wb.create_sheet('Truy vet');trace.append(['Yêu cầu','Số TC','ID'])
    for fr in [f'FR-{i:02}' for i in range(1,13)]+['BR/Auth','BR/Security','UI']:
        selected=[r[0] for r in CASES if fr in r[2]]; trace.append([fr,len(selected),', '.join(selected)])
    style(trace,[20,15,130],100)
    questions=wb.create_sheet('Can lam ro');questions.append(['ID','Chủ đề','Phát hiện','Xử lý'])
    for r in QUESTIONS:questions.append(r)
    style(questions,[15,35,85,100],110)
    verification=wb.create_sheet('Kiem chung');verification.append(['Hạng mục','Kết quả / giới hạn'])
    for r in [('SRS','Đã đọc bản gốc v1.0; có nhiều TBD và mâu thuẫn quyền Admin/history.'),('Build','PASS npm run build; warning chunk-size không chặn build.'),('Tests có sẵn','19/19 PASS; không phải toàn bộ bộ TC mới.'),('Browser có sẵn','PASS reader borrow, admin approve, CRUD book, users/profile/mobile; no pageerror.'),('Checks bổ trợ','43/43 assertions đạt; 33 registration, 8 invalid-date, 2 quan sát API. Hai quan sát không phải kết luận web đạt.'),('Lỗi web xác nhận','BUG-01: email83 ký tự đăng nhập API200; formmaxlength64 cắt còn64, UI401. TC BB2-222 kỳ vọng thành công, chưa sửa ứng dụng.'),('Sai lệch nhãn','BUG-02: newest có nhãn Mới cập nhật nhưng sắp id DESC; chờ chốt định nghĩa theo Q-01.'),('Giới hạn','Chưa chạy toàn bộ TC, rate-limit100/15 và phiên8giờ chưa chạy trong review. Không tự ghi Pass vào bộ mẫu. Bằng chứng trong evidence/ và srs-code-review.md.')]:verification.append(r)
    style(verification,[30,145],85)
    wb.save(ROOT/'test-cases.xlsx')
    shutil.copyfile(ROOT/'test-cases.xlsx', ROOT.parent.parent/'TEST_CASE_HOP_DEN.xlsx')
    loaded=load_workbook(ROOT/'test-cases.xlsx');assert loaded['Test cases'].max_row==len(CASES)+1
    assert (ROOT/'test-cases.xlsx').read_bytes()==(ROOT.parent.parent/'TEST_CASE_HOP_DEN.xlsx').read_bytes()
    md=['# Kiểm thử hộp đen — Lá Thư Viện, phiên bản 2.1','',f'**{len(CASES)} test case riêng biệt**, tất cả chưa chạy. Mỗi biến thể đầu vào có ID và ô kết quả riêng. Một TC có thể có nhiều bước đọc để kiểm tra cùng một hành vi.','', '[Tải Excel](test-cases.xlsx) · [CSV](test-cases.csv) · [Script sinh tài liệu](generate_workbook.py)','', 'Nguồn: [OpenAPI](../openapi.json), [quyết định](../decisions.md), [truy vết](../traceability.md), [README](../../README.md). Đã đối chiếu SRS gốc v1.0 do người dùng gửi; xem [báo cáo SRS/code](srs-code-review.md). Đối chiếu mã để tìm khoảng trống/mâu thuẫn, không lấy hành vi triển khai làm chuẩn thay SRS.','', '## Hướng dẫn','']
    for key,value in GUIDE:md += [f'**{key}:** {value}','']
    md += ['## Điểm cần làm rõ','', '| ID | Chủ đề | Phát hiện | Xử lý |','| --- | --- | --- | --- |']
    for r in QUESTIONS:md.append('| '+' | '.join(r)+' |')
    md += ['', '## Test case','']
    for r in CASES:
        md += [f'### {r[0]} — {r[3]}', '',f'- Truy vết: {r[2]}; ID cũ: {r[1]}; ưu tiên: {r[4]}; kỹ thuật: {r[5]}.',f'- Nguồn kỳ vọng: {r[6]}.', f'- Căn cứ SRS/phạm vi: {alignment(r)[0]}', f'- Vị trí code đối chiếu tĩnh: {alignment(r)[1]}',f'- Tiền điều kiện: {r[7]}.',f'- Dữ liệu: {r[8]}.','',r[9],'',f'**HTTP mong đợi:** {r[10]}.',f'**Kết quả mong đợi:** {r[11]}',f'**Hậu điều kiện/dọn dữ liệu:** {r[12]}','**Trạng thái:** Chưa chạy.','']
    (ROOT/'test-plan.md').write_text('\n'.join(md),encoding='utf-8')
    report='''# Rà soát thiết kế test case — phiên bản 2.0

Đã sửa gộp biến thể, bổ sung tiền điều kiện/fixture và dữ liệu nền, tách các hành vi độc lập, chuẩn hóa HTTP/field/status và kiểm tra trạng thái trước/sau. CSV và hai bản Excel được sinh từ cùng một nguồn; không có kết quả thực thi giả định.

- Tách chuỗi/độ dài, trường thiếu, ngày, phân trang, vai trò, trạng thái và quyết định thành từng TC.
- Thay chuỗi tìm rỗng có wildcard bằng token chữ/số; ghi rõ vai trò tạo fixture và vai trò gọi request.
- Bổ sung phân trang liên tiếp/trang cuối/vượt trang, DENY khi hết tồn, tổng bằng số đang mượn, ghi đè availableQuantity, xử lý ID không tồn tại, gửi lại sau APPROVED và DENIED, API quản trị bị chặn theo từng operation.
- Tách thống kê thành ba biến đổi có delta rõ; PUT giữ các trường không thử, tránh vô tình xóa email.
- Tách responsive theo vai trò/viewport; tách SQL input và HTML; ghi điều kiện rate-limit và session để tránh false pass.
- Ghi riêng các khoảng trống đặc tả, nhất là sort=newest, query ngoài hợp đồng và ADMIN cuối. Không tự xác nhận yêu cầu chưa chốt.

## Kiểm tra tài liệu đã thực hiện

Kiểm tra ID duy nhất, đủ các cột thiết kế bắt buộc, trạng thái đều Chưa chạy, ánh xạ toàn bộ 64 ID cũ, truy vết FR-01…12, đọc lại XLSX bằng openpyxl, CSV round-trip và hai file XLSX giống nhau. Đây là kiểm tra cấu trúc tài liệu, không phải chạy TC trên ứng dụng.
'''
    (ROOT/'review-v2.md').write_text(report,encoding='utf-8')
    print('Generated',len(CASES),'cases; priorities',dict(Counter(r[4] for r in CASES)))

if __name__=='__main__':main()
