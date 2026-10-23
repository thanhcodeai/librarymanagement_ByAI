"""Build an execution report only from a completed 30-case evidence JSON."""
import argparse
import csv
import hashlib
import json
import shutil
import struct
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
REPO=HERE.parent.parent
GITHUB='https://github.com/thanhcodeai/librarymanagement_ByAI/blob/blackbox-test-cases/'

def local(iso):
    return datetime.fromisoformat(iso.replace('Z','+00:00')).astimezone(ZoneInfo('Asia/Saigon')).strftime('%d/%m/%Y %H:%M:%S')

def style(ws,widths,height=110):
    ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    for i,v in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=v
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='215E78');c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height=height
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        if row[0].row%2==0:
            for c in row:c.fill=PatternFill('solid',fgColor='EEF4FA')


def main():
    ap=argparse.ArgumentParser();ap.add_argument('evidence');args=ap.parse_args()
    evidence=Path(args.evidence).resolve();data=json.loads(evidence.read_text());results=data['results']
    assert len(results)==30 and {r['id'] for r in results}=={f'TC-{i:02}' for i in range(1,31)}, 'Need all 30 unique results, not a partial/zero-test run'
    with (HERE/'simple-test-cases.csv').open(encoding='utf-8-sig') as f:plan={r['ID']:r for r in csv.DictReader(f)}
    for r in results:
        assert r['expected']==plan[r['id']]['Kết quả mong đợi'], 'Expected outcome must not be rewritten to fit observations'
        for name in r['screenshots']:assert (evidence.parent/name).is_file()
        assert r['status'] in ['Pass','Fail','Blocked']
    counts=Counter(r['status'] for r in results);passed=counts['Pass'];failed=counts['Fail'];blocked=counts['Blocked']
    start=min(r['start'] for r in results);finish=max(r['finish'] for r in results)
    pct=round(passed/30*100,2);elapsed=(datetime.fromisoformat(finish.replace('Z','+00:00'))-datetime.fromisoformat(start.replace('Z','+00:00'))).total_seconds()
    outdir=evidence.parent;relative=outdir.relative_to(REPO).as_posix()
    headers=['ID','SRS','Test case','Kết quả mong đợi','Kết quả thực tế','Kết luận','HTTP quan sát','Thời gian (Asia/Saigon)','Bằng chứng','Ghi chú']
    rows=[]
    for r in results:
        http='; '.join(dict.fromkeys(f"{x['method']} {x['path']} → {x['status']}" for x in r['http']))
        if r.get('apiEvidence'):http='GET /api/v1/admin/stats → '+str(r['apiEvidence']['status'])
        proofs=', '.join(r['screenshots']) or 'results.json → '+r['id']+' → apiEvidence'
        actual=r['actual'] or 'Không hoàn thành: '+r['error']
        note=plan[r['id']]['Ghi chú / bằng chứng']
        if r.get('retestNote'):note+='\n'+r['retestNote']
        if r['error']:note+='\n'+r['error']
        rows.append([r['id'],r['requirement'],r['name'],r['expected'],actual,r['status'],http,local(r['start'])+' – '+local(r['finish']),proofs,note.strip()])
    summary=[
        ('Hệ thống','Lá Thư Viện'),('Loại kiểm thử','Hộp đen chức năng: thao tác qua Chromium cho29 ca;1 ca HTTP kiểm tra quyền. Playwright mô phỏng các bước của bộ TC thủ công.'),
        ('Bộ TC','30 ca trong docs/blackbox/simple-test-cases.csv; không phải bộ226 ca.'),('SRS','Bản gốc v1.0 ngày16/10/2026 do người dùng gửi; xem srs-code-review.md về các TBD.'),
        ('Mã ứng dụng',data['applicationCommit']),('Run ID',data['run']),('Bắt đầu / kết thúc',local(start)+' – '+local(finish)+' (Asia/Saigon)'),('Thời lượng',f'{elapsed:.2f} giây (không tính cài đặt/dựng database ban đầu)'),
        ('Môi trường',f"Node {data['node']}; Chromium {data['browserVersion']}; MySQL8.4; viewport1440×1060; database riêng {data['database']}"),
        ('Dữ liệu','Fixture riêng theo từng ca, dựng qua API công khai trước bước thử; tác vụ mục tiêu29 ca đi qua UI. Không dùng hoặc sửa dữ liệu vận hành. Hồ sơ thử có lịch sử được giữ trong DB test để truy vết.'),
        ('Tổng / thực hiện',f'30 / {len(results)}'),('Pass',str(passed)),('Fail',str(failed)),('Blocked',str(blocked)),('Chưa chạy','0'),('Tỷ lệ đạt',f'{pct}% trên30 ca đã chọn'),
        ('Bằng chứng',relative+'/results.json và ảnh PNG từng ca; TC-30 có body HTTP đã loại thông tin xác thực.'),
        ('Lưu ý HTTP','400/401/403/409/404 xuất hiện ở ca âm có thể là kết quả mong đợi; không tự coi mọi HTTP4xx là test Fail.'),
        ('Lượt chốt','Báo cáo tổng hợp1 lượt đủ30 ca và1 lượt chạy lại riêngTC-18/TC-27. Kết quả gốc từng lượt được giữ trong sourceRuns của results.json; ghi chú retest có ở hai dòng tương ứng. Các lượt chuẩn bị có lỗi locator/thiết lập kịch bản đã được chẩn đoán và chạy lại sau sửa runner; không tính lỗi công cụ thành lỗi website. Không sửa mã ứng dụng hoặc đổi kỳ vọng để làm test đạt.'),
        ('Lỗi ngoài bộ30','BUG-01 đã xác nhận ở lượt đối chiếu trước: email hợp lệ83 ký tự bị form cắt còn64; API login200 nhưng UI401. Chưa sửa; không nằm trong30 ca và không tính thêm vào mẫu số30.'),
        ('Điểm cần chốt','UX-02: nhãn Mới cập nhật dùng thứ tự tạo; các TBD của SRS về password/session/email/Guest/giới hạn mượn vẫn cần quyết định bổ sung.'),
        ('Chưa bao phủ','Không thực thi toàn bộ226 ca, mọi biên/kiểu dữ liệu, đồng thời, fault-injection rollback, chờ8giờ, rate-limit, load/backup/restore, productionHTTPS. Không coi pass30 là đạt toàn bộ SRS hay hết lỗi.'),
        ('Kết luận','Các chức năng cơ bản trong30 ca đạt trên môi trường thử.' if passed==30 else f'Cần xử lý {failed} ca Fail và {blocked} ca Blocked trước khi kết luận bộ30 ca đạt.'),
    ]
    wb=Workbook();ws=wb.active;ws.title='Tong hop';ws.append(['Nội dung','Kết quả'])
    for r in summary:ws.append(r)
    style(ws,[30,145],80);ws.freeze_panes='A2'
    detail=wb.create_sheet('Ket qua 30 TC');detail.append(headers)
    for row in rows:detail.append(row)
    style(detail,[12,24,38,70,85,15,85,42,42,70],135)
    for i,row in enumerate(rows,2):
        detail.cell(i,6).font=Font(bold=True,color={'Pass':'008000','Fail':'C00000','Blocked':'9C6500'}[row[5]])
        original=results[i-2]
        target=relative+'/'+original['screenshots'][0] if original['screenshots'] else relative+'/results.json'
        detail.cell(i,9).hyperlink=GITHUB+target;detail.cell(i,9).font=Font(color='0563C1',underline='single')
    issues=wb.create_sheet('Luu y va loi da biet');issues.append(['ID','Tình trạng','Nội dung','Xử lý / phạm vi'])
    issues.append(['BUG-01','Mở; ngoài30 ca','Email83 ký tự: API login200, form cắt64 và UI401.','Đã có evidence/ui-review.json từ lượt đối chiếu trước; sửa form, rồi retest TC email trong bộ chi tiết.'])
    issues.append(['UX-02','Cần chốt đặc tả','Nhãn Mới cập nhật nhưng sort theo ID tạo.','Chốt ý nghĩa trước khi sửa UI hoặc API; chưa kết luận vi phạm một yêu cầu sort cụ thể của SRS.'])
    issues.append(['SRS-TBD','Chưa nghiệm thu','Các chính sách còn TBD trong SRS gốc.','Phê duyệt/cập nhật baseline, không tự coi quyết định repo là mọi yêu cầu SRS đã chốt.'])
    style(issues,[18,32,90,100],110)
    xlsx=outdir/'KET_QUA_KIEM_THU_HOP_DEN.xlsx';wb.save(xlsx);shutil.copyfile(xlsx,REPO/xlsx.name)
    with (outdir/'ket-qua-30-tc.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f,lineterminator='\n');w.writerow(headers);w.writerows(rows)
    md=['# Kết quả kiểm thử hộp đen — Lá Thư Viện','',f'**Run:** `{data["run"]}`. **Kết quả:** {passed} Pass, {failed} Fail, {blocked} Blocked, 0 chưa chạy; tỷ lệ đạt {pct}% trong30 ca đã chọn.','',
        '## Môi trường, phương pháp và phạm vi','']
    for key,value in summary:md.extend([f'- **{key}:** {value}'])
    md+=['','## Kết quả từng ca','','| ID | SRS | Nội dung | Kết quả thực tế | Kết luận | Bằng chứng |','| --- | --- | --- | --- | --- | --- |']
    for row,r in zip(rows,results):
        proof=f'[{r["screenshots"][0]}]({r["screenshots"][0]})' if r['screenshots'] else '[HTTP JSON](results.json)'
        cells=[row[0],row[1],row[2],row[4].replace('\n','<br>'),row[5],proof]
        md.append('| '+' | '.join(v.replace('|','\\|') for v in cells)+' |')
    md+=['','## Tài liệu đính kèm','','- [Excel kết quả](KET_QUA_KIEM_THU_HOP_DEN.xlsx)','- [Báo cáo Word](BAO_CAO_KIEM_THU_HOP_DEN.docx)','- [Bằng chứng JSON](results.json)','- [CSV](ket-qua-30-tc.csv)','',
         'Ảnh được chụp sau bước xác minh của từng ca; riêngTC-15 có hai phiên người dùng, mỗi phiên một ảnh. Mật khẩu/cookie không được lưu trong báo cáo. Dữ liệu nhận dạng trong ảnh là tài khoản thử.']
    (outdir/'bao-cao-ket-qua.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    doc=Document();sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.6)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    doc.add_heading('BÁO CÁO KẾT QUẢ KIỂM THỬ HỘP ĐEN',0);doc.add_paragraph('Website quản lý thư viện — Lá Thư Viện')
    doc.add_paragraph(f'Lượt chạy: {data["run"]}\nThời gian: {local(start)} – {local(finish)} (Asia/Saigon)')
    doc.add_heading('1. Mục tiêu và phạm vi',1)
    doc.add_paragraph('Kiểm tra30 test case cơ bản đã thống nhất, dựa trên SRS gốc và giao diện hiện có. Đây là kết quả thực thi bằng Playwright trên Chromium (29 ca UI) và HTTP client (1 ca API), không phải kết quả suy luận từ code. Fixture được tạo bằng API trước bước kiểm thử; không truy vấn DB để quyết định Pass/Fail.')
    doc.add_heading('2. Môi trường',1)
    for key in ['SRS','Mã ứng dụng','Môi trường','Dữ liệu']:
        doc.add_paragraph(key+': '+dict(summary)[key])
    doc.add_heading('3. Kết quả tổng hợp',1)
    t=doc.add_table(rows=1,cols=6);t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,['Tổng','Pass','Fail','Blocked','Chưa chạy','Tỷ lệ đạt']):c.text=v
    for c,v in zip(t.add_row().cells,['30',str(passed),str(failed),str(blocked),'0',str(pct)+'%']):c.text=v
    doc.add_paragraph('Tỷ lệ chỉ áp dụng cho30 ca trong lượt chạy này. Các lỗi input bị chặn đúng như yêu cầu (HTTP4xx) được đánh giá Pass khi phản hồi và trạng thái dữ liệu đúng kỳ vọng.')
    doc.add_heading('4. Kết quả chi tiết',1)
    table=doc.add_table(rows=1,cols=6);table.style='Table Grid'
    for c,v in zip(table.rows[0].cells,['ID','SRS','Test case','Mong đợi','Thực tế','Kết luận']):c.text=v
    trpr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for row in rows:
        for cell,value in zip(table.add_row().cells,[row[0],row[1],row[2],row[3],row[4],row[5]]):cell.text=value
    table.autofit=False
    for row in table.rows:
        for c,width in zip(row.cells,[.55,.95,1.6,3.0,3.0,.75]):c.width=Inches(width)
        for c in row.cells:
            for para in c.paragraphs:
                for run_ in para.runs:run_.font.size=Pt(9)
    doc.add_heading('5. Lỗi đã biết và giới hạn',1)
    for key in ['Lỗi ngoài bộ30','Điểm cần chốt','Chưa bao phủ','Lượt chốt','Kết luận']:doc.add_paragraph(dict(summary)[key])
    doc.add_heading('6. Một số bằng chứng',1)
    for ident in ['TC-02','TC-12','TC-24','TC-26']:
        r=next(x for x in results if x['id']==ident)
        if r['screenshots']:
            doc.add_paragraph(ident+' — '+r['name']);image_path=outdir/r['screenshots'][0]; iw,ih=struct.unpack('>II',image_path.read_bytes()[16:24]); scale=min(8.6/iw,5.2/ih); doc.add_picture(str(image_path),width=Inches(iw*scale),height=Inches(ih*scale))
    doc.add_paragraph('Toàn bộ ảnh, HTTP metadata và dữ liệu quan sát: '+GITHUB+relative+'/results.json')
    docx=outdir/'BAO_CAO_KIEM_THU_HOP_DEN.docx';doc.save(docx);shutil.copyfile(docx,REPO/docx.name)
    manifest={'run':data['run'],'counts':dict(counts),'total':30,'percentPass':pct,'startLocal':local(start),'finishLocal':local(finish),'rawEvidenceSHA256':hashlib.sha256(evidence.read_bytes()).hexdigest(),'runnerSHA256':hashlib.sha256((HERE/'run-simple-blackbox.mjs').read_bytes()).hexdigest(),'planSHA256':hashlib.sha256((HERE/'simple-test-cases.csv').read_bytes()).hexdigest(),'report':'bao-cao-ket-qua.md','applicationCommit':data['applicationCommit']}
    (outdir/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    assert load_workbook(xlsx)['Ket qua 30 TC'].max_row==31
    assert len(Document(docx).tables[1].rows)==31
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
