"""Generate the short manual test set; preserve the full suite in test-cases.xlsx."""
from pathlib import Path
import csv
import shutil
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent
# Requirement, name, preconditions/data, steps, expected outcome, note
CASES=[
('FR-01','Đăng ký hợp lệ','Chưa đăng nhập. Username=readerRUN chưa tồn tại; họ tên=Nguyễn An; mật khẩu thử=Abcd1234. Email có thể để trống theo cấu hình web.', '1. Bấm Đăng ký độc giả.\n2. Nhập dữ liệu trên, bấm Tạo tài khoản.\n3. Đăng nhập tài khoản vừa tạo.', 'Thông báo đăng ký thành công; đăng nhập được và vào Khám phá sách với quyền độc giả.', 'Email tùy chọn và mật khẩu tối thiểu8 thuộc quyết định triển khai, SRS còn TBD.'),
('FR-01','Bỏ trống trường bắt buộc','Chưa đăng nhập; họ tên=Nguyễn An, mật khẩu thử=Abcd1234; bỏ trống tên đăng nhập.', '1. Mở Đăng ký độc giả.\n2. Điền họ tên/mật khẩu, để trống tên đăng nhập.\n3. Bấm Tạo tài khoản.', 'Không tạo tài khoản; form chỉ ra tên đăng nhập bắt buộc; không vào màn hình đã đăng nhập.', ''),
('FR-01','Tên đăng nhập đã tồn tại','Chưa đăng nhập; đã có tài khoản readerRUN.', '1. Mở Đăng ký độc giả.\n2. Dùng username readerRUN, họ tên/mật khẩu hợp lệ.\n3. Bấm Tạo tài khoản.', 'Hiện lỗi thông tin đã tồn tại; không tạo tài khoản thứ hai và không đổi tài khoản cũ.', ''),
('FR-02','Đăng nhập độc giả','U1 là tài khoản USER có mật khẩu thử đã biết; hiện chưa đăng nhập.', '1. Nhập username/mật khẩu đúng của U1.\n2. Bấm Đăng nhập.', 'Vào Khám phá sách; có Sách của tôi và Hồ sơ cá nhân; hiển thị đúng người đăng nhập.', ''),
('FR-02','Đăng nhập quản trị viên','A1 là tài khoản ADMIN có mật khẩu thử đã biết; hiện chưa đăng nhập.', '1. Nhập username/mật khẩu đúng của A1.\n2. Bấm Đăng nhập.', 'Vào Tổng quan thư viện; có Quản lý sách, Yêu cầu mượn, Người dùng.', ''),
('FR-02','Đăng nhập sai mật khẩu','Chưa đăng nhập; U1 tồn tại.', '1. Nhập username U1, mật khẩu sai.\n2. Bấm Đăng nhập.', 'Báo tên đăng nhập hoặc mật khẩu không đúng; vẫn ở màn hình đăng nhập.', ''),
('FR-03','Tìm thấy sách','Đăng nhập U1; có sách tên=KiemThuRUN thuộc bộ dữ liệu thử.', '1. Ở Khám phá sách, nhập KiemThuRUN vào ô Tìm sách.\n2. Chờ danh sách cập nhật.', 'Danh sách có sách KiemThuRUN; các kết quả khớp từ khóa trong tên/tác giả/mã sách.', ''),
('FR-03','Tìm không có kết quả','Đăng nhập U1; không có sách chứa token chữ/số ZZZRUNNoMatch999.', '1. Tìm ZZZRUNNoMatch999.\n2. Chờ danh sách cập nhật.', 'Hiện trạng thái không có kết quả; không báo lỗi hệ thống và không hiển thị sách không khớp.', 'Không dùng % hoặc _ vì web hiện coi chúng là wildcard.'),
('FR-03','Lọc thể loại','Đăng nhập U1. Có B1 thể loại CAT-RUN-A và B2 thể loại CAT-RUN-B.', '1. Xóa từ khóa tìm kiếm cũ.\n2. Chọn thể loại CAT-RUN-A.', 'Có B1, không có B2; mọi sách trong danh sách thuộc CAT-RUN-A.', ''),
('FR-03','Chuyển trang danh sách','Đăng nhập U1; có13 sách thử trong thể loại PAGE-RUN, không có sách khác thuộc thể loại này.', '1. Lọc PAGE-RUN.\n2. Ghi ID/mã sách trang1.\n3. Bấm trang kế tiếp.\n4. Quay lại trang1.', 'Trang1 có12 sách, trang2 có1 sách; không trùng/mất sách giữa hai trang; quay lại đúng danh sách ban đầu.', '12 sách/trang là cấu hình web hiện tại, không phải con số đã chốt trong SRS.'),
('FR-04','Xem chi tiết sách','Đăng nhập U1; B1 có tên, tác giả, mã, thể loại, mô tả và tồn đã biết.', '1. Tìm B1.\n2. Bấm xem sách.', 'Mở Chi tiết sách; tên/tác giả/mã/thể loại/mô tả và số có sẵn/tổng khớp dữ liệu B1.', ''),
('FR-05','Gửi yêu cầu mượn hợp lệ','Đăng nhập U1; B1 có sẵn2 cuốn; U1 chưa có yêu cầu chờ duyệt B1. D là ngày mai, D+1 là ngày kế tiếp.', '1. Mở B1, bấm Gửi yêu cầu mượn.\n2. Nhập ngày mượn D, ngày trả D+1.\n3. Gửi; mở Sách của tôi.\n4. Mở lại chi tiết B1.', 'Báo gửi thành công; có yêu cầu mới Chờ duyệt với đúng sách/ngày; B1 vẫn có sẵn2 cuốn.', ''),
('FR-05 / BR-05','Ngày trả bằng ngày mượn','Đăng nhập U1; B1 tồn tại; chưa có yêu cầu chờ duyệt cùng sách. D là ngày mai.', '1. Mở form mượn B1.\n2. Đặt cả ngày mượn và ngày trả là D.\n3. Bấm Gửi yêu cầu; kiểm tra Sách của tôi.', 'Không tạo yêu cầu; báo ngày trả phải sau ngày mượn; tồn sách không đổi.', 'Chọn một giá trị đại diện nhóm ngày sai; các biên khác ở bộ chi tiết.'),
('FR-06','Xem hồ sơ cá nhân','Đăng nhập U1; biết họ tên, username, email và role của U1.', '1. Mở Hồ sơ cá nhân.\n2. So sánh thông tin với tài khoản U1.', 'Đúng hồ sơ U1; không hiển thị mật khẩu hoặc password hash.', ''),
('FR-07 / BR-03','Chỉ xem lịch sử của mình','U1 có yêu cầu R1, U2 có yêu cầu R2 khác nhau. Biết mã yêu cầu; dùng hai phiên trình duyệt riêng.', '1. Đăng nhập U1, mở Sách của tôi, xem đủ các trang.\n2. Ở phiên khác đăng nhập U2, làm tương tự.', 'U1 thấy R1, không thấy R2; U2 thấy R2, không thấy R1; mỗi danh sách chỉ có yêu cầu của chính người đăng nhập.', 'Đây là quan sát UI; kiểm tra chống giả mạo API đầy đủ vẫn ở bộ chi tiết.'),
('FR-07','Lọc lịch sử mượn','Đăng nhập U1; có ít nhất một yêu cầu Chờ duyệt và một yêu cầu Đã duyệt.', '1. Mở Sách của tôi.\n2. Chọn tab Chờ duyệt.', 'Chỉ hiện yêu cầu Chờ duyệt của U1; không có yêu cầu đã duyệt/từ chối.', ''),
('FR-08','Đăng xuất','Đăng nhập U1.', '1. Bấm Đăng xuất.\n2. Tải lại trang.', 'Trở về màn hình đăng nhập; không tự vào lại tài khoản U1.', 'Ca cơ bản UI; kiểm tra gửi lại cookie cũ thuộc bộ chi tiết.'),
('FR-09','Thống kê cập nhật sau thêm sách','Đăng nhập A1; ghi số Đầu sách=N và tổng cuốn=M trên Tổng quan. Không có người khác thay đổi dữ liệu.', '1. Vào Quản lý sách, thêm sách mới mã STATS-RUN, tổng2 cuốn; các trường bắt buộc hợp lệ.\n2. Quay lại Tổng quan.', 'Đầu sách=N+1; tổng cuốn=M+2; không tự tăng số Chờ duyệt/Đã duyệt.', ''),
('FR-10','Thêm sách hợp lệ','Đăng nhập A1; mã BK-TC19-RUN chưa tồn tại. Tên=Sach Thu RUN; tác giả=Tac Gia; thể loại=CAT-RUN; tổng3.', '1. Mở Quản lý sách, bấm Thêm sách mới.\n2. Điền dữ liệu; bấm Lưu thay đổi.\n3. Tìm mã sách vừa tạo.', 'Tạo thành công; đúng thông tin đã nhập; tổng3 và có sẵn3.', 'Tồn khởi tạo bằng tổng là cách web hiện triển khai.'),
('FR-10','Không thêm trùng mã sách','Đăng nhập A1; đã có sách mã BK-DUP-RUN.', '1. Thêm sách mới dùng mã BK-DUP-RUN.\n2. Các trường còn lại hợp lệ; bấm Lưu thay đổi.', 'Báo mã sách đã tồn tại; không tạo bản sao, không đổi sách cũ.', ''),
('FR-10','Sửa thông tin sách','Đăng nhập A1; B1 có tổng3/có sẵn3, chưa có yêu cầu mượn.', '1. Trong Quản lý sách, bấm Sửa B1.\n2. Đổi tên thành Ten Moi RUN, giữ các trường khác.\n3. Lưu; tìm và mở lại B1.', 'Tên mới được lưu; mã/tác giả/tổng/có sẵn giữ nguyên.', ''),
('FR-10','Xóa sách chưa có lịch sử','Đăng nhập A1; B1 là sách thử không có bất kỳ yêu cầu mượn nào.', '1. Bấm Xóa B1.\n2. Xác nhận xóa.\n3. Tìm lại đúng mã sách.', 'Xóa thành công; không còn B1 trong danh sách/kết quả tìm.', ''),
('FR-10 / BR-15','Không xóa sách có lịch sử','Đăng nhập A1; B1 có ít nhất một yêu cầu đã từ chối, chưa bị xóa.', '1. Bấm Xóa B1.\n2. Xác nhận.\n3. Đọc lại sách và yêu cầu liên quan.', 'Báo không thể xóa do có lịch sử; sách và yêu cầu cũ vẫn còn.', ''),
('FR-11 / BR-08','Duyệt yêu cầu khi còn sách','Đăng nhập A1; R1 đang Chờ duyệt cho B1; B1 có sẵn2 cuốn.', '1. Vào Yêu cầu mượn, tìm R1.\n2. Bấm Duyệt, ghi chú Dong y RUN; xác nhận.\n3. Kiểm tra yêu cầu và số có sẵn của B1.', 'R1 thành Đã duyệt, ghi chú được lưu; B1 có sẵn1 cuốn; tổng sách không đổi.', ''),
('FR-11 / BR-09','Từ chối yêu cầu','Đăng nhập A1; R1 đang Chờ duyệt; B1 có sẵn2 cuốn.', '1. Tìm R1 trong Yêu cầu mượn.\n2. Bấm Từ chối, ghi chú Ly do RUN; xác nhận.\n3. Kiểm tra R1 và B1.', 'R1 thành Từ chối, lưu ghi chú; B1 vẫn có sẵn2 cuốn.', ''),
('FR-11 / BR-10','Duyệt khi hết sách','A1/U1/U2 có tài khoản. B1 ban đầu có1 cuốn; U1/U2 đã gửi R1/R2; A1 đã duyệt R2 nên B1 có sẵn0, R1 còn Chờ duyệt.', '1. A1 tìm R1 trong Yêu cầu mượn.\n2. Bấm Duyệt và xác nhận.\n3. Kiểm tra R1 và B1.', 'Báo sách đã hết; R1 vẫn Chờ duyệt; B1 có sẵn0, không âm.', 'Chuẩn bị dữ liệu qua web theo tiền điều kiện, không chỉnh trực tiếp database.'),
('FR-12','Thêm người dùng','Đăng nhập A1; username memberRUN chưa tồn tại; họ tên=Tran B; mật khẩu thử=Abcd1234; role=Độc giả.', '1. Vào Người dùng, bấm Thêm người dùng.\n2. Điền dữ liệu, lưu.\n3. Tìm lại username memberRUN.', 'Có tài khoản mới, đúng họ tên/username và vai trò Độc giả; không hiển thị mật khẩu.', ''),
('FR-12','Sửa người dùng','Đăng nhập A1; U2 là độc giả thử có thông tin đã biết.', '1. Trong Người dùng, bấm Sửa U2.\n2. Đổi họ tên thành Ten Moi RUN, giữ username/email/role, để trống mật khẩu mới.\n3. Lưu và mở lại.', 'Họ tên đổi đúng; username/email/role giữ nguyên.', ''),
('FR-12','Xóa người dùng chưa có lịch sử','Đăng nhập A1; U2 là độc giả thử không có yêu cầu mượn, không phải A1.', '1. Trong Người dùng, bấm Xóa U2.\n2. Xác nhận.\n3. Tìm lại username U2.', 'Xóa thành công; danh sách không còn U2.', ''),
('BR-04 / NFR-SEC-03','Độc giả không gọi được API quản trị','Có phiên đăng nhập U1 trong Postman hoặc HTTP client. Đây là ca API duy nhất của bộ rút gọn.', '1. Bằng cookie phiên U1, gọi GET /api/v1/admin/stats.\n2. Đọc HTTP và body, không dùng phiên ADMIN.', 'HTTP403; success=false; không trả số liệu thống kê quản trị.', 'Ẩn menu trên UI chưa đủ chứng minh phân quyền backend. Dùng client hỗ trợ cookie; không ghi cookie vào báo cáo.')
]
GUIDE=[
('Bộ kiểm thử','30 test case cơ bản: 29 ca thao tác web, 1 ca API kiểm tra phân quyền. Phủ 12 nhóm FR; là bộ kiểm tra chức năng chính, không tuyên bố nghiệm thu đầy đủ mọi SRS/NFR.'),
('Căn cứ','SRS gốc v1.0 ngày16/10/2026 do người dùng gửi; đối chiếu src/main.jsx, server/app.js và quyết định trong docs/decisions.md. Báo cáo chi tiết: docs/blackbox/srs-code-review.md.'),
('Nguồn SRS','Đặc tả yêu cầu phần mềm website quản lý thư viện — SRS.docx; SHA-256 a440fa65c73192e27e82b27dee6a8bbef78505d31ae98023388250db3356e703.'),
('Chuẩn bị','Dùng môi trường test, không dùng production. A1=ADMIN, U1/U2=hai USER riêng. Dùng username/mật khẩu thực đã có, không đăng nội dung .env, mật khẩu hoặc cookie vào báo cáo/Git.'),
('Dữ liệu','RUN là mã chữ/số riêng cho lượt chạy, ví dụ20261023A; thay RUN trong dữ liệu trước khi nhập. B1/R1 là tên gọi đối tượng, lấy đúng sách/yêu cầu đã chuẩn bị; mỗi ca chuẩn bị theo tiền điều kiện, không phụ thuộc kết quả ca trước.'),
('Tạo dữ liệu','A1 tạo sách và người dùng qua màn hình quản trị. USER tạo yêu cầu; ADMIN duyệt/từ chối để dựng lịch sử. Giữ nguyên số liệu đúng tiền điều kiện trước ca; dùng sách thử riêng cho ca thay đổi tồn.'),
('Mật khẩu ví dụ','Abcd1234 chỉ là dữ liệu cho tài khoản thử tạo mới, không phải mật khẩu mặc định của tài khoản seed hay tài khoản thật.'),
('Cách ghi kết quả','Pass: làm đủ bước và mọi kết quả đúng. Fail: có sai khác. Blocked: thiếu dữ liệu/công cụ. Chưa chạy: chưa thực hiện. Điền cột Kết quả thực tế và Ghi chú/bằng chứng; không tự lấy kết quả test cũ làm Pass.'),
('Trạng thái bộ mới','Tất cả Chưa chạy. Việc bộ chi tiết từng có kiểm tra bổ trợ không có nghĩa 30 ca này đã được thực hiện đầy đủ.'),
('Điểm SRS còn TBD','Auth/session/timeout, password, email optional, quyền Guest, duplicatePENDING, giới hạn mượn và quyền ADMIN khác vẫn TBD trong SRS. Ca có phụ thuộc cấu hình hiện tại được ghi ở cột Ghi chú, không gọi là yêu cầu đã chốt của SRS.'),
('Lỗi web đã biết','FR-02: email hợp lệ83 ký tự đăng nhập API được nhưng ô đăng nhập UI chỉ giữ64 ký tự. Lỗi vẫn chưa sửa. Bộ rút gọn dùng username cho ca cơ bản; lỗi email vẫn ghi nhận tại đây, không bị xóa khỏi báo cáo chi tiết.'),
('Sắp xếp chưa chốt','Nhãn Mới cập nhật hiện sắp theo ID tạo mới. Vì SRS chưa chốt tạo mới hay lần chỉnh sửa, không đưa cách hiểu này làm tiêu chí Pass của bộ cơ bản.'),
('Không có trong 30 ca','Biên độ dài/kiểu dữ liệu chi tiết, API malformed/not-found, fake session/spoof, đồng thời, rollback, replay cookie, rate-limit, chờ8giờ, mọi tổ hợp CRUD/role. Tham khảo bộ226 ca và báo cáo nếu cần kiểm tra sâu. Không kết luận toàn bộ các phần này đạt.'),
('Ngoài phạm vi','Không thêm trả sách, gia hạn, phạt, thanh toán vì SRS chưa đưa vào phạm vi. Hiệu năng, SLA, HTTPS production, backup/RPO/RTO không được nghiệm thu bằng bộ cơ bản này.'),
('Giữ bản chi tiết','Bộ226 ca vẫn ở docs/blackbox/test-cases.xlsx. File TEST_CASE_HOP_DEN.xlsx tại thư mục gốc bây giờ là bản30 ca đơn giản.'),
('Dọn dữ liệu','Chỉ xóa tài khoản/sách thử chưa có lịch sử bằng UI sau khi lưu bằng chứng. Dữ liệu đã có lịch sử giữ trong môi trường test; không xóa database/volume để dọn tùy tiện.')
]
HEAD=['ID','SRS','Test case','Tiền điều kiện / dữ liệu','Bước thực hiện','Kết quả mong đợi','Trạng thái','Kết quả thực tế','Ghi chú / bằng chứng']

def format_sheet(ws,widths,height):
    ws.freeze_panes='D2' if ws.title=='Test case' else 'A2'
    ws.auto_filter.ref=ws.dimensions
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='215E78');c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=30
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height=height
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if row[0].row%2==0:c.fill=PatternFill('solid',fgColor='EEF4FA')
    ws.sheet_view.zoomScale=85

def main():
    assert len(CASES)==30
    rows=[[f'TC-{i:02}',fr,name,pre,steps,expect,'Chưa chạy','',note] for i,(fr,name,pre,steps,expect,note) in enumerate(CASES,1)]
    wb=Workbook();ws=wb.active;ws.title='Test case';ws.append(HEAD)
    for r in rows:ws.append(r)
    format_sheet(ws,[12,24,36,62,70,70,18,40,65],125)
    v=DataValidation(type='list',formula1='"Chưa chạy,Pass,Fail,Blocked"');v.showErrorMessage=True;v.error='Chọn trạng thái trong danh sách.';v.errorStyle='stop';ws.add_data_validation(v);v.add('G2:G31')
    guide=wb.create_sheet('Huong dan');guide.append(['Mục','Nội dung'])
    for row in GUIDE:guide.append(row)
    format_sheet(guide,[30,140],85)
    wb.save(ROOT/'simple-test-cases.xlsx');shutil.copyfile(ROOT/'simple-test-cases.xlsx',ROOT.parent.parent/'TEST_CASE_HOP_DEN.xlsx')
    with (ROOT/'simple-test-cases.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f,lineterminator='\n');w.writerow(HEAD);w.writerows(rows)
    md=['# Bộ test case cơ bản — 30 ca','', 'Ưu tiên thao tác thủ công trên web, bám SRS gốc và màn hình đã code. Đây là bộ kiểm tra rút gọn, không thay bộ nghiệm thu đầy đủ.','', '[Excel30 ca](simple-test-cases.xlsx) · [CSV](simple-test-cases.csv) · [Bộ chi tiết226 ca](test-cases.xlsx) · [Báo cáo SRS/code](srs-code-review.md)','', '## Hướng dẫn','']
    for title,text in GUIDE:md.extend([f'**{title}:** {text}',''])
    md.extend(['## Test case',''])
    for r in rows:
        md.extend([f'### {r[0]} — {r[2]}',f'**SRS:** {r[1]}',f'**Tiền điều kiện/dữ liệu:** {r[3]}','',r[4],'',f'**Mong đợi:** {r[5]}','**Trạng thái:** Chưa chạy.',f'**Ghi chú:** {r[8]}' if r[8] else '', ''])
    (ROOT/'simple-test-plan.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    assert load_workbook(ROOT/'simple-test-cases.xlsx')['Test case'].max_row==31
    print('Generated 30 cases, 2 sheets. Full 226-case suite retained separately.')
if __name__=='__main__':main()
